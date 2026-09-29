from io import BytesIO

import pytest
from docx import Document
from docx.oxml import OxmlElement

from app.temporary.parsing import parse_temporary


def test_plain_and_markdown_preserve_the_entire_text_and_code_indentation():
    text = "# Heading\n\n```py\ndef solve():\n    return 42\n```\nlast line\n"
    assert parse_temporary("a.md", text.encode()) == text
    assert parse_temporary("a.py", text.encode()) == text


def test_word_preserves_heading_tables_and_order():
    doc = Document()
    doc.add_heading("Heading", level=1)
    doc.add_paragraph("before")
    doc.add_table(rows=1, cols=2).rows[0].cells[0].text = "table content"
    doc.add_paragraph("after")
    buffer = BytesIO()
    doc.save(buffer)
    text = parse_temporary("a.docx", buffer.getvalue())
    assert text.index("Heading") < text.index("before") < text.index("table content") < text.index("after")


def test_invalid_encoding_and_pdf_are_rejected_instead_of_silent_loss():
    with pytest.raises(ValueError): parse_temporary("a.txt", b"\xff\xfe\x00")
    with pytest.raises(ValueError): parse_temporary("a.pdf", b"data")


def test_word_embedded_content_is_rejected_instead_of_silently_ignored():
    doc = Document()
    doc.add_paragraph("visible text")._p.append(OxmlElement("w:drawing"))
    buffer = BytesIO()
    doc.save(buffer)
    with pytest.raises(ValueError, match="TEMPORARY_UNSUPPORTED_EMBEDDED_CONTENT"):
        parse_temporary("a.docx", buffer.getvalue())
